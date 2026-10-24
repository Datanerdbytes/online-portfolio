"""Synthetic documents and provider failures; no real personal data or API calls."""

import io
import json
import zipfile

import httpx
import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.services import resume_extraction as extraction
from app.services.resume_extraction import ImportProblem, extract_document, validate_extraction

TEXT = "Alex Example\nData Engineer\nBuilds reliable data pipelines with Python."


def pdf(pages=None, encrypted=False):
    writer = PdfWriter()
    for value in pages or [TEXT]:
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        stream = DecodedStreamObject()
        stream.set_data(
            (
                "BT /F1 12 Tf 50 700 Td (" + value.replace("\n", ") Tj 0 -18 Td (") + ") Tj ET"
            ).encode()
        )
        page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("test-only")
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def docx(text=TEXT):
    document = Document()
    document.add_paragraph(text)
    row = document.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text = "Python"
    row.cells[1].text = "SQL"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def proposal():
    return {
        "candidates": [
            {
                "section": "profile",
                "values": {"name": "Alex Example"},
                "evidence": {"name": "Alex Example"},
                "generated": [],
                "warnings": [],
            }
        ],
        "not_imported": [],
    }


def test_reads_pdf_and_word_tables():
    assert "Alex Example" in extract_document(pdf(), ".pdf")
    assert "Python | SQL" in extract_document(docx(), ".docx")


@pytest.mark.parametrize(
    "data,extension,match",
    [
        (b"fake.pdf", ".pdf", "content"),
        (b"%PDF-corrupt", ".pdf", "could not read"),
        (pdf(encrypted=True), ".pdf", "Password"),
        (pdf([""]), ".pdf", "readable"),
        (pdf([TEXT, ""]), ".pdf", "readable"),
        (pdf([TEXT] * 11), ".pdf", "10 pages"),
        (b"x" * (3 * 1024 * 1024 + 1), ".docx", "3 MiB"),
    ],
)
def test_rejects_bad_and_incomplete_documents(data, extension, match):
    with pytest.raises(ImportProblem, match=match):
        extract_document(data, extension)


def test_rejects_expanded_word_and_too_much_text():
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"x" * (21 * 1024 * 1024))
    with pytest.raises(ImportProblem, match="complex"):
        extract_document(output.getvalue(), ".docx")
    with pytest.raises(ImportProblem, match="50,000"):
        extract_document(docx("x" * 50001), ".docx")


def test_excerpts_and_generated_fields_are_verified():
    assert validate_extraction(proposal(), TEXT).candidates[0].values["name"] == "Alex Example"
    invented = proposal()
    invented["candidates"][0]["evidence"]["name"] = "Invented Person"
    corrected = validate_extraction(invented, TEXT)
    assert corrected.candidates[0].evidence["name"] == "Alex Example"
    invented["candidates"][0]["values"]["name"] = "Unsupported person"
    with pytest.raises(ImportProblem, match="verified"):
        validate_extraction(invented, TEXT)
    generated = proposal()
    generated["candidates"][0]["generated"] = ["name"]
    with pytest.raises(ImportProblem, match="verified"):
        validate_extraction(generated, TEXT)
    generated = proposal()
    generated["candidates"][0]["values"]["headline"] = "Data professional"
    generated["candidates"][0]["generated"] = ["headline"]
    assert validate_extraction(generated, TEXT).candidates[0].generated == ["headline"]


@pytest.mark.parametrize(
    "status,finish,content",
    [(429, "stop", "{}"), (200, "length", "{}"), (200, "stop", "not json"), (200, "stop", "{}")],
)
def test_provider_failures_do_not_return_candidates(monkeypatch, status, finish, content):
    original = httpx.AsyncClient

    def handler(request):
        return httpx.Response(
            status, json={"choices": [{"finish_reason": finish, "message": {"content": content}}]}
        )

    monkeypatch.setattr(
        extraction.httpx,
        "AsyncClient",
        lambda **kw: original(transport=httpx.MockTransport(handler), **kw),
    )
    with pytest.raises(ImportProblem):
        extraction.map_resume(TEXT)


def test_provider_receives_only_source_and_mapping_instructions(monkeypatch):
    original = httpx.AsyncClient

    def handler(request):
        body = json.loads(request.content)
        assert body["messages"][1] == {"role": "user", "content": TEXT}
        assert "tools" not in body and body["stream"] is False
        assert body["response_format"]["type"] == "json_schema"
        assert body["chat_template_kwargs"]["enable_thinking"] is False
        return httpx.Response(
            200,
            json={
                "choices": [
                    {"finish_reason": "stop", "message": {"content": json.dumps(proposal())}}
                ]
            },
        )

    monkeypatch.setattr(
        extraction.httpx,
        "AsyncClient",
        lambda **kw: original(transport=httpx.MockTransport(handler), **kw),
    )
    assert extraction.map_resume(TEXT).candidates


def test_real_isolated_parser_worker():
    assert "Alex Example" in extraction.extract_bounded(docx(), ".docx")
    with pytest.raises(ImportProblem, match="content"):
        extraction.extract_bounded(b"not a pdf", ".pdf")


def test_normalizes_document_ligatures_without_generating_words():
    text = extract_document(
        docx("Alex Example uses Snowﬂake for reliable analytics pipelines."), ".docx"
    )
    assert "Snowflake" in text
    assert "Alex Example" in text


def test_absent_evidence_null_is_not_accepted_as_proof():
    payload = proposal()
    payload["candidates"][0]["values"]["location"] = None
    payload["candidates"][0]["evidence"]["location"] = None
    assert "location" not in validate_extraction(payload, TEXT).candidates[0].evidence
    payload["candidates"][0]["evidence"]["name"] = None
    payload["candidates"][0]["values"]["name"] = "Unsupported person"
    with pytest.raises(ImportProblem, match="verified"):
        validate_extraction(payload, TEXT)


def test_missing_list_null_is_absence_not_a_factual_value():
    payload = proposal()
    payload["candidates"].append(
        {
            "section": "experience",
            "values": {"role": "Data Engineer", "highlights": None},
            "evidence": {"role": "Data Engineer", "highlights": None},
            "generated": [],
            "warnings": [],
        }
    )
    result = validate_extraction(payload, TEXT)
    assert "highlights" not in result.candidates[1].values


def test_missing_scalar_citation_requires_the_value_to_exist_in_source():
    payload = proposal()
    payload["candidates"][0]["evidence"] = {}
    assert validate_extraction(payload, TEXT).candidates[0].evidence["name"] == "Alex Example"
    payload["candidates"][0]["values"]["name"] = "Invented Name"
    with pytest.raises(ImportProblem, match="verified"):
        validate_extraction(payload, TEXT)


def test_reflowed_quote_is_replaced_with_actual_source_span():
    payload = proposal()
    payload["candidates"][0]["evidence"]["name"] = "Alex Example"
    result = validate_extraction(payload, "Alex\nExample: Data Engineer")
    assert result.candidates[0].evidence["name"] == "Alex\nExample"


def test_empty_suggestion_is_omitted_instead_of_inventing_content():
    payload = proposal()
    payload["candidates"][0]["values"]["headline"] = None
    payload["candidates"][0]["generated"] = ["headline"]
    candidate = validate_extraction(payload, TEXT).candidates[0]
    assert "headline" not in candidate.values
    assert candidate.generated == []


def test_list_items_are_grounded_individually_and_unverified_items_omitted():
    payload = proposal()
    payload["candidates"].append(
        {
            "section": "skills",
            "values": {"label": "Data Engineer", "items": ["Data Engineer", "Unsupported tool"]},
            "evidence": {"label": "Data Engineer", "items": "invented joined quote"},
            "generated": [],
            "warnings": [],
        }
    )
    candidate = validate_extraction(payload, TEXT).candidates[1]
    assert candidate.values["items"] == ["Data Engineer"]
    assert candidate.evidence["items"] == ["Data Engineer"]
    assert candidate.warnings


def test_wrapped_url_is_verified_without_following_it():
    payload = proposal()
    payload["candidates"].append(
        {
            "section": "links",
            "values": {"label": "GitHub", "url": "https://github.com/example"},
            "evidence": {"label": "GitHub", "url": "https://github.com/example"},
            "generated": [],
            "warnings": [],
        }
    )
    result = validate_extraction(payload, TEXT + "\nGitHub: https://github.com/\nexample")
    assert result.candidates[1].evidence["url"] == "https://github.com/\nexample"
